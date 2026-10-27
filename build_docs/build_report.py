"""Builds docs/IDS_CASE_STUDY_REPORT.pdf and docs/IDS_URLs_conclusion.docx.
Reads results/metrics.json and figures/*.png when present; otherwise results sections
carry a clearly marked 'pending' notice.   Usage: python3 build_docs/build_report.py
"""
import json, os
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                PageBreak, Image, ListFlowable, ListItem, KeepTogether)
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "docs"); os.makedirs(OUT, exist_ok=True)
mp = os.path.join(ROOT, "results", "metrics.json")
M = json.load(open(mp)) if os.path.exists(mp) else None
def fig(n):
    p = os.path.join(ROOT, "figures", n); return p if os.path.exists(p) else None
def f(v, d=3):
    return "-" if v is None else f"{float(v):.{d}f}"

NAME, ROLL, BRANCH, SEC = "A Badri", "160124771305", "AI & DS", "I1"
TITLE = "IPL Score Forecasting & Batting-Collapse Detection using Machine Learning in R"
GREEN, RED = colors.HexColor("#0B3D2E"), colors.HexColor("#E4572E")

ss = getSampleStyleSheet()
H1 = ParagraphStyle("H1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=16, textColor=GREEN, spaceBefore=14, spaceAfter=6, keepWithNext=1)
H2 = ParagraphStyle("H2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=12.5, textColor=GREEN, spaceBefore=8, spaceAfter=4, keepWithNext=1)
B = ParagraphStyle("B", parent=ss["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=6)
CAP = ParagraphStyle("CAP", parent=B, fontSize=9, textColor=colors.HexColor("#6B7280"), alignment=1)
NOTE = ParagraphStyle("NOTE", parent=B, backColor=colors.HexColor("#FFF4E5"), borderPadding=6, textColor=colors.HexColor("#7A4B00"))

def bullets(items):
    return ListFlowable([ListItem(Paragraph(i, B), leftIndent=12) for i in items], bulletType="bullet", start="-", leftIndent=14)

def table(rows, widths=None, hi_row=None):
    t = Table(rows, colWidths=widths, repeatRows=1)
    st = [("BACKGROUND", (0, 0), (-1, 0), GREEN), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
          ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9.5),
          ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
          ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F0")]),
          ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
    if hi_row: st.append(("BACKGROUND", (0, hi_row), (-1, hi_row), colors.HexColor("#F2C14E")))
    t.setStyle(TableStyle(st)); return t

def figure(name, caption, w=14):
    p = fig(name)
    if not p:
        return [Paragraph(f"[Figure pending: {caption} - created when IDS_CS.Rmd is knitted on the data]", NOTE)]
    from PIL import Image as PI
    iw, ih = PI.open(p).size
    return [KeepTogether([Image(p, width=w * cm, height=w * cm * ih / iw), Paragraph(caption, CAP)])]

def footer(c, d):
    c.saveState(); c.setFont("Helvetica", 8.5); c.setFillColor(colors.HexColor("#6B7280"))
    c.drawString(2 * cm, 1.2 * cm, f"{NAME} | {ROLL} | {BRANCH} | Section {SEC}")
    c.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {d.page}"); c.restoreState()

S = []
# ---- cover
S += [Spacer(1, 4 * cm), Paragraph("INTRODUCTION TO DATA SCIENCE - CASE STUDY REPORT", ParagraphStyle("c0", parent=B, textColor=RED, fontName="Helvetica-Bold", fontSize=11)),
      Paragraph(TITLE, ParagraphStyle("c1", parent=H1, fontSize=24, leading=30, spaceBefore=10)), Spacer(1, 1 * cm),
      table([["Name", NAME], ["Roll No", ROLL], ["Branch", BRANCH], ["Section", SEC], ["Dataset", "Cricsheet IPL ball-by-ball JSON, 1,169 matches"], ["Language", "R (R Markdown)"]],
            [4 * cm, 11 * cm]), PageBreak()]

# ---- abstract & intro
S += [Paragraph("1. Abstract", H1),
      Paragraph("This case study builds a data-science pipeline in R on 1,169 official, ball-by-ball IPL matches. Using the runs and wickets of the "
                "last 12 balls together with the current match state, it (i) forecasts the final innings score at 6, 10, 15 and 18 overs and "
                "(ii) estimates the probability that the batting side will collapse, defined as losing three or more wickets in the next 12 legal balls. "
                "Several regression and classification models are compared on matches never seen during training.", B),
      Paragraph("2. Introduction and Problem Statement", H1),
      Paragraph("A run-rate projection (current run rate x 20) is the usual way to estimate a T20 total, but it cannot tell a side that is "
                "scoring steadily with wickets in hand from one that has just lost two quick wickets. Collapses are rare but decisive events, so "
                "an early warning is valuable to analysts and captains. The project therefore asks two questions:", B),
      bullets(["What is the expected final score of the innings at a given checkpoint?",
               "How likely is a batting collapse in the next 12 balls?"]),
      Paragraph("3. Objectives", H1),
      bullets(["Parse official ball-by-ball JSON into a clean, analysis-ready table.",
               "Engineer momentum and pressure features from the last 12 balls.",
               "Compare baseline, linear, ridge, random-forest and gradient-boosting forecasters.",
               "Define a collapse mathematically and train classifiers to detect it.",
               "Provide a reproducible R Markdown report and an interactive what-if simulator."])]

# ---- data
S += [Paragraph("4. Dataset", H1),
      Paragraph("Data come from Cricsheet (cricsheet.org), which publishes official ball-by-ball records, one JSON file per match. "
                "Each file contains an <b>info</b> block (teams, venue, season, overs) and an <b>innings</b> list; each innings contains "
                "<b>overs</b>, and each over contains <b>deliveries</b> with <i>runs</i> (batter, extras, total), optional <i>extras</i> "
                "(wides, no-balls) and optional <i>wickets</i>.", B),
      Paragraph("4.1 Cleaning rules", H2),
      bullets(["Only innings 1 and 2 of full 20-over matches are kept (super overs and rain-shortened games removed).",
               "Wides and no-balls are folded into the legal ball they accompany so that the ball counter is exact.",
               "'Retired hurt' is not counted as a wicket.",
               "Every legal ball becomes one row with score, wickets, runs and wickets in the last 12 balls, and wickets in the next 12 balls."])]
if M:
    S.append(Paragraph(f"After cleaning: <b>{M.get('n_matches','-')}</b> matches, <b>{M.get('n_innings','-')}</b> innings and <b>{M.get('n_balls','-')}</b> legal-ball rows.", B))
else:
    S.append(Paragraph("[Pending: exact counts after cleaning are written to results/metrics.json when the Rmd is knitted.]", NOTE))

# ---- method
S += [Paragraph("5. Methodology", H1),
      Paragraph("5.1 Features", H2),
      table([["Feature", "Meaning"], ["cum_score, cum_wickets", "Score and wickets at the checkpoint"], ["current_run_rate", "Score divided by overs bowled"],
             ["runs_last12", "Runs in the last 12 legal balls (momentum)"], ["wkts_last12", "Wickets in the last 12 legal balls (pressure)"],
             ["checkpoint_overs", "6, 10, 15 or 18 overs completed"]], [5 * cm, 11 * cm]),
      Paragraph("5.2 Collapse definition", H2),
      Paragraph("A collapse starts at ball b if the batting side loses <b>3 or more wickets in the next 12 legal balls</b>. Inputs come only from balls "
                "already bowled, so there is no information leakage. Windows that run past an all-out innings are padded with 'no more wickets'; "
                "windows cut short by a successful chase are discarded.", B),
      Paragraph("5.3 Models", H2),
      bullets(["<b>Forecasting:</b> mean baseline, run-rate x 20, linear regression, ridge regression (glmnet), random forest, gradient boosting (gbm, trees chosen by 5-fold CV).",
               "<b>Collapse detection:</b> logistic regression and random-forest classifier; decision threshold chosen to maximise F1 on training data only."]),
      Paragraph("5.4 Validation", H2),
      Paragraph("Matches are split 80/20 at random (seed 123) so that no match contributes to both sets. Forecasts are scored with MAE, RMSE and R-squared; "
                "collapse classifiers with AUC, precision, recall, F1 and accuracy, since collapses are a minority class.", B)]

# ---- results
S += [Paragraph("6. Results", H1)]
if M:
    S.append(Paragraph("6.1 Score forecasting", H2))
    rows = [["Model", "MAE", "RMSE", "R-squared"]]
    hi = None
    for i, r in enumerate(M.get("forecast", []), 1):
        rows.append([r["Model"], f(r["MAE"], 2), f(r["RMSE"], 2), f(r["R2"], 3)])
        if r["Model"] == M.get("best_model"): hi = i
    S += [table(rows, [6 * cm, 3 * cm, 3 * cm, 3 * cm], hi), Paragraph("Hold-out matches; best model by MAE highlighted.", CAP)]
    c = M.get("collapse", {})
    S.append(Paragraph("6.2 Collapse detection", H2))
    rows = [["Model", "AUC", "Precision", "Recall", "F1", "Accuracy"]]
    for nm, k in (("Logistic regression", "logistic"), ("Random forest", "random_forest")):
        r = c.get(k, {}); rows.append([nm, f(r.get("auc")), f(r.get("precision")), f(r.get("recall")), f(r.get("f1")), f(r.get("accuracy"))])
    S += [table(rows, [4.6 * cm, 2.2 * cm, 2.4 * cm, 2.2 * cm, 2.2 * cm, 2.4 * cm]),
          Paragraph(f"Collapse base rate in the test matches: {f(c.get('base_rate'), 3)}.", CAP)]
else:
    S.append(Paragraph("[Pending] Results tables are filled automatically from results/metrics.json after the Rmd is knitted on the full dataset "
                       "and 'python3 build_docs/build_report.py' is re-run. No figures have been invented for this draft.", NOTE))
S += figure("fig1_score_distribution.png", "Figure 1: Distribution of first-innings totals", 12)
S += figure("fig3_mae_by_checkpoint.png", "Figure 2: Forecast error by checkpoint", 12)
S += figure("fig4_feature_importance.png", "Figure 3: Random-forest feature importance", 12)
S += figure("fig5_roc.png", "Figure 4: ROC curves for collapse detection", 12)
S += figure("fig6_collapse_heatmap.png", "Figure 5: Collapse rate by overs and wickets down", 12)

S += [Paragraph("7. Discussion", H1),
      Paragraph("The forecasting models are compared against the run-rate x 20 projection at each checkpoint, so the benefit of the "
                "last-12-ball features can be read directly from the tables above. Because collapses are a minority class, "
                "accuracy alone is not informative; AUC, precision and recall describe how useful the early warning really is.", B),
      Paragraph("8. Limitations and Future Work", H1),
      bullets(["No pitch, weather, opposition or batter-quality features.",
               "Only first-innings totals are forecast because chasing innings end early.",
               "Random match split; a season-wise split would test drift over time.",
               "The collapse rule (3 wickets in 12 balls) is a modelling choice; other windows can be tried.",
               "Future work: batter/bowler form, venue effects, calibrated boosted classifiers, a hosted dashboard."]),
      Paragraph("9. Conclusion", H1),
      Paragraph("A complete, reproducible R pipeline turns raw official IPL ball-by-ball JSON into two practical outputs: an expected final score "
                "and a short-horizon collapse probability, both driven by the runs and wickets of the last 12 balls and the current match state.", B),
      Paragraph("10. References", H1),
      bullets(["Cricsheet - ball-by-ball data and JSON format: https://cricsheet.org/downloads/ and https://cricsheet.org/format/json/",
               "Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.",
               "Friedman, J. H. (2001). Greedy function approximation: a gradient boosting machine. Annals of Statistics, 29(5).",
               "Friedman, Hastie, Tibshirani (2010). Regularization paths for generalized linear models via coordinate descent. J. Stat. Software 33(1).",
               "R packages: jsonlite, dplyr, ggplot2, caret, randomForest, glmnet, gbm, pROC (CRAN)."])]

SimpleDocTemplate(os.path.join(OUT, "IDS_CASE_STUDY_REPORT.pdf"), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                  topMargin=2 * cm, bottomMargin=2 * cm, title=TITLE, author=NAME).build(S, onFirstPage=footer, onLaterPages=footer)

# ================= URLs & conclusion docx =================
d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
h = d.add_heading("IDS Case Study - URLs and Conclusion", 0)
p = d.add_paragraph(); r = p.add_run(f"{NAME}  |  Roll No: {ROLL}  |  Branch: {BRANCH}  |  Section: {SEC}"); r.bold = True
d.add_paragraph("Project: " + TITLE)
d.add_heading("Project URLs", 1)
urls = [("GitHub repository (replace with your link)", "https://github.com/<your-username>/ipl-collapse-forecast"),
        ("Dataset - Cricsheet IPL downloads", "https://cricsheet.org/downloads/"),
        ("Cricsheet JSON format documentation", "https://cricsheet.org/format/json/"),
        ("R project", "https://www.r-project.org/"),
        ("randomForest (CRAN)", "https://cran.r-project.org/package=randomForest"),
        ("glmnet (CRAN)", "https://cran.r-project.org/package=glmnet"),
        ("gbm (CRAN)", "https://cran.r-project.org/package=gbm"),
        ("pROC (CRAN)", "https://cran.r-project.org/package=pROC"),
        ("ggplot2 (CRAN)", "https://cran.r-project.org/package=ggplot2")]
t = d.add_table(rows=1, cols=2); t.style = "Light Grid Accent 1"
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Resource", "URL"
for a, b in urls:
    row = t.add_row().cells; row[0].text, row[1].text = a, b
d.add_heading("Conclusion", 1)
for para in [
    "This case study applied a complete data-science workflow in R to 1,169 official ball-by-ball IPL matches: parsing nested JSON, engineering features, training and comparing models, and validating them on unseen matches.",
    "Two outputs were produced from the runs and wickets of the last 12 balls and the current match state: (1) the expected final innings score at 6, 10, 15 and 18 overs and (2) the probability of a batting collapse, defined as three or more wickets in the next 12 legal balls.",
    "Simple baselines (mean and run-rate x 20) were compared with linear, ridge, random-forest and gradient-boosting regressors, and collapse detection was evaluated with AUC, precision, recall and F1 because collapses are rare.",
]:
    d.add_paragraph(para)
if M:
    best = M.get("best_model"); fr = {r["Model"]: r for r in M.get("forecast", [])}
    c = M.get("collapse", {})
    d.add_paragraph(f"Key results on hold-out matches: best forecaster = {best} (MAE {f(fr.get(best, {}).get('MAE'), 2)} runs vs {f(fr.get('Run-rate x 20', {}).get('MAE'), 2)} for run-rate x 20); "
                    f"best collapse AUC = {f(max(c.get('logistic', {}).get('auc', 0), c.get('random_forest', {}).get('auc', 0)))}.")
else:
    pr = d.add_paragraph(); rr = pr.add_run("[Add your key numbers here after knitting IDS_CS.Rmd - MAE of the best model vs run-rate x 20, and the collapse AUC / F1. Or re-run build_docs/build_report.py and they are inserted automatically.]")
    rr.italic = True; rr.font.color.rgb = RGBColor(0xB4, 0x5A, 0x00)
d.add_paragraph("Limitations: no pitch, weather or player-quality features; first-innings forecasting only; random match split. Future work: batter/bowler form, venue effects, season-wise validation and a hosted dashboard.")
d.save(os.path.join(OUT, "IDS_URLs_conclusion.docx"))
print("Wrote report PDF and URLs docx", "(with metrics)" if M else "(results pending)")
